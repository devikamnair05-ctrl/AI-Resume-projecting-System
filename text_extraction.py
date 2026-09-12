"""
text_extraction.py
Module for extracting text and metadata from PDF, DOCX, and TXT resumes.
Supports both file paths on disk and in-memory file streams (e.g. Streamlit UploadedFile).
"""

import io
import os
import re
from typing import Union, Tuple, Dict, Any, Optional

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

try:
    import docx
except ImportError:
    docx = None


def extract_text_from_pdf(file_input: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extract text from a PDF file using PyMuPDF.
    
    Args:
        file_input: File path (str) or binary content (bytes/BytesIO).
        
    Returns:
        Extracted text string.
    """
    if fitz is None:
        raise ImportError("PyMuPDF is not installed. Please install it with 'pip install PyMuPDF'.")
        
    text_parts = []
    
    if isinstance(file_input, str):
        doc = fitz.open(file_input)
    elif isinstance(file_input, bytes):
        doc = fitz.open(stream=file_input, filetype="pdf")
    elif isinstance(file_input, io.BytesIO):
        doc = fitz.open(stream=file_input.getvalue(), filetype="pdf")
    else:
        # Handles Streamlit UploadedFile (has .read() / .getvalue())
        if hasattr(file_input, "getvalue"):
            doc = fitz.open(stream=file_input.getvalue(), filetype="pdf")
        elif hasattr(file_input, "read"):
            doc = fitz.open(stream=file_input.read(), filetype="pdf")
        else:
            raise TypeError(f"Unsupported file_input type for PDF extraction: {type(file_input)}")
            
    try:
        for page_num in range(len(doc)):
            page = doc[page_num]
            page_text = page.get_text()
            if page_text:
                text_parts.append(page_text)
    finally:
        doc.close()
        
    return "\n".join(text_parts).strip()


def extract_text_from_docx(file_input: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extract text from a Word (.docx) document using python-docx.
    
    Args:
        file_input: File path (str) or binary content (bytes/BytesIO).
        
    Returns:
        Extracted text string.
    """
    if docx is None:
        raise ImportError("python-docx is not installed. Please install it with 'pip install python-docx'.")
        
    if isinstance(file_input, str):
        doc = docx.Document(file_input)
    elif isinstance(file_input, bytes):
        doc = docx.Document(io.BytesIO(file_input))
    elif isinstance(file_input, io.BytesIO):
        doc = docx.Document(file_input)
    else:
        if hasattr(file_input, "getvalue"):
            doc = docx.Document(io.BytesIO(file_input.getvalue()))
        elif hasattr(file_input, "read"):
            doc = docx.Document(io.BytesIO(file_input.read()))
        else:
            raise TypeError(f"Unsupported file_input type for DOCX extraction: {type(file_input)}")
            
    text_parts = []
    # Extract from paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())
            
    # Extract from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)
                
    return "\n".join(text_parts).strip()


def extract_text_from_txt(file_input: Union[str, bytes, io.BytesIO]) -> str:
    """Extract text from plain text file or stream."""
    if isinstance(file_input, str):
        with open(file_input, "r", encoding="utf-8", errors="replace") as f:
            return f.read().strip()
    elif isinstance(file_input, bytes):
        return file_input.decode("utf-8", errors="replace").strip()
    elif isinstance(file_input, io.BytesIO):
        return file_input.getvalue().decode("utf-8", errors="replace").strip()
    elif hasattr(file_input, "getvalue"):
        return file_input.getvalue().decode("utf-8", errors="replace").strip()
    elif hasattr(file_input, "read"):
        content = file_input.read()
        if isinstance(content, bytes):
            return content.decode("utf-8", errors="replace").strip()
        return str(content).strip()
    return ""


def extract_text(file_input: Union[str, bytes, io.BytesIO], filename: Optional[str] = None) -> str:
    """
    Universal text extraction dispatcher based on file extension.
    
    Args:
        file_input: File path, bytes, or Streamlit UploadedFile.
        filename: Optional filename to identify extension if file_input is a stream.
        
    Returns:
        Extracted text.
    """
    ext = ""
    if filename:
        _, ext = os.path.splitext(filename)
    elif isinstance(file_input, str):
        _, ext = os.path.splitext(file_input)
    elif hasattr(file_input, "name"):
        _, ext = os.path.splitext(file_input.name)
        
    ext = ext.lower()
    
    if ext == ".pdf":
        return extract_text_from_pdf(file_input)
    elif ext in {".docx", ".doc"}:
        return extract_text_from_docx(file_input)
    elif ext == ".txt":
        return extract_text_from_txt(file_input)
    else:
        # Fallback attempt PDF first, then TXT
        try:
            return extract_text_from_pdf(file_input)
        except Exception:
            return extract_text_from_txt(file_input)


def extract_metadata(text: str, filename: str = "") -> Dict[str, Any]:
    """
    Extract basic metadata such as candidate name, email, and phone heuristics from resume text.
    """
    metadata = {
        "candidate_name": "Unknown Candidate",
        "email": None,
        "phone": None
    }
    
    # Heuristic 1: Extract email
    email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
    emails = re.findall(email_pattern, text)
    if emails:
        metadata["email"] = emails[0]
        
    # Heuristic 2: Extract phone
    phone_pattern = r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}'
    phones = re.findall(phone_pattern, text)
    if phones:
        # Join groups if tuple
        match = re.search(phone_pattern, text)
        if match:
            metadata["phone"] = match.group(0).strip()
            
    # Heuristic 3: Extract name from first non-empty lines
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    candidate_name = None
    
    for line in lines[:5]:
        # Exclude lines that look like emails, urls, or section headers
        clean_line = re.sub(r'[^a-zA-Z\s]', '', line).strip()
        words = clean_line.split()
        if 2 <= len(words) <= 4 and not any(k in line.lower() for k in ["resume", "curriculum", "email", "phone", "summary", "experience"]):
            candidate_name = clean_line.title()
            break
            
    if not candidate_name and filename:
        # Derive name from filename (e.g. resume1_anu_ml_senior.pdf -> Anu Ml Senior)
        base = os.path.splitext(os.path.basename(filename))[0]
        cleaned_base = re.sub(r'^(resume\d*|cv\d*)_?', '', base, flags=re.IGNORECASE)
        candidate_name = cleaned_base.replace("_", " ").replace("-", " ").strip().title()
        
    metadata["candidate_name"] = candidate_name if candidate_name else "Candidate"
    return metadata
